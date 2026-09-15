from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "IMPLEMENTACION_BACKUP_SUPABASE.docx"


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_code(document, code):
    paragraph = document.add_paragraph()
    paragraph.style = document.styles["No Spacing"]
    paragraph.paragraph_format.left_indent = Inches(0.35)
    paragraph.paragraph_format.right_indent = Inches(0.35)
    run = paragraph.add_run(code)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(45, 45, 45)
    return paragraph


def add_bullet(document, text, level=0):
    paragraph = document.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.add_run(text)
    return paragraph


document = Document()
section = document.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

styles = document.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(28)
styles["Title"].font.color.rgb = RGBColor(62, 92, 78)
styles["Heading 1"].font.name = "Aptos Display"
styles["Heading 1"].font.color.rgb = RGBColor(62, 92, 78)
styles["Heading 2"].font.name = "Aptos Display"
styles["Heading 2"].font.color.rgb = RGBColor(169, 93, 48)

header = section.header.paragraphs[0]
header.text = "LAS FLORES | DOCUMENTO TÉCNICO"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor(110, 110, 110)

footer = section.footer.paragraphs[0]
footer.text = "Implementación de backup de Supabase | 14 de septiembre de 2026"
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.runs[0].font.size = Pt(8)
footer.runs[0].font.color.rgb = RGBColor(110, 110, 110)

title = document.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Implementación del Backup\nde Supabase")
subtitle = document.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("Restaurante Las Flores")
run.bold = True
run.font.size = Pt(15)
run.font.color.rgb = RGBColor(169, 93, 48)

document.add_paragraph()
summary = document.add_paragraph()
summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
summary.add_run(
    "Documento técnico y operativo para respaldar diariamente la base de datos PostgreSQL "
    "de Supabase mediante GitHub Actions."
)

table = document.add_table(rows=4, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Light Shading Accent 1"
metadata = [
    ("Proyecto", "Las Flores Elevated Web App"),
    ("Fecha", "14 de septiembre de 2026"),
    ("Responsable técnico", "Equipo de desarrollo"),
    ("Estado", "Implementado en la rama feature/supabase-backup"),
]
for row, (key, value) in zip(table.rows, metadata):
    set_cell_text(row.cells[0], key, bold=True, color=(62, 92, 78))
    set_cell_text(row.cells[1], value)

document.add_page_break()

document.add_heading("1. Resumen ejecutivo", level=1)
document.add_paragraph(
    "Se implementó un mecanismo de backup lógico automatizado para la base de datos "
    "de Las Flores. Cada día, GitHub Actions instala las herramientas de PostgreSQL, "
    "ejecuta pg_dump contra Supabase usando una conexión TLS y publica el resultado "
    "como un artifact privado con una retención de 30 días."
)
document.add_paragraph(
    "El diseño evita almacenar contraseñas, URLs de conexión o archivos de respaldo "
    "en el repositorio. También incorpora un checksum SHA-256 para verificar que el "
    "archivo descargado no fue alterado antes de restaurarlo."
)

document.add_heading("2. Objetivos y alcance", level=1)
document.add_heading("Objetivos", level=2)
for item in [
    "Proteger los datos operativos del restaurante ante borrados, errores o incidentes.",
    "Automatizar el backup diario y permitir ejecuciones manuales.",
    "Permitir la restauración controlada mediante pg_restore.",
    "Mantener las credenciales fuera del código fuente.",
]:
    add_bullet(document, item)
document.add_heading("Qué incluye", level=2)
for item in [
    "Esquema public de PostgreSQL.",
    "Tablas, datos, índices, funciones, triggers y políticas RLS.",
    "Archivo en formato custom de PostgreSQL.",
    "Archivo de checksum SHA-256.",
]:
    add_bullet(document, item)
document.add_heading("Qué no incluye", level=2)
for item in [
    "Objetos binarios de Supabase Storage, como imágenes y archivos.",
    "El esquema administrado auth y los usuarios gestionados por Supabase.",
    "Una restauración automática sobre producción.",
]:
    add_bullet(document, item)

document.add_heading("3. Arquitectura de la solución", level=1)
document.add_paragraph("El flujo implementado es:")
add_code(document, "GitHub Actions -> pg_dump -> archivo .dump + checksum -> artifact privado")
document.add_paragraph(
    "El workflow se ejecuta diariamente a las 05:00 UTC y también puede iniciarse "
    "desde la opción Run workflow de GitHub. El secreto SUPABASE_DB_URL se inyecta "
    "solo durante la ejecución y no aparece en los logs ni en los archivos generados."
)

architecture = document.add_table(rows=1, cols=3)
architecture.alignment = WD_TABLE_ALIGNMENT.CENTER
architecture.style = "Light Shading Accent 1"
for cell, text in zip(architecture.rows[0].cells, ["Componente", "Responsabilidad", "Ubicación"]):
    set_cell_text(cell, text, bold=True, color=(255, 255, 255))
    set_cell_shading(cell, "3E5C4E")
rows = [
    ("GitHub Actions", "Programa y ejecuta el backup", ".github/workflows/supabase-backup.yml"),
    ("PowerShell", "Valida la conexión y ejecuta pg_dump", "scripts/backup-supabase.ps1"),
    ("Supabase PostgreSQL", "Origen de los datos", "Proyecto Supabase configurado"),
    ("GitHub Artifact", "Almacenamiento temporal privado", "Retención de 30 días"),
]
for values in rows:
    cells = architecture.add_row().cells
    for cell, text in zip(cells, values):
        set_cell_text(cell, text)

document.add_heading("4. Archivos implementados", level=1)
files = [
    (".github/workflows/supabase-backup.yml", "Programa el backup, instala PostgreSQL y sube el artifact."),
    ("scripts/backup-supabase.ps1", "Valida SUPABASE_DB_URL, exige TLS, ejecuta pg_dump y crea el checksum."),
    ("scripts/backup-supabase.test.ps1", "Comprueba la configuración, el modo dry-run y la sintaxis PowerShell."),
    ("docs/SUPABASE_BACKUP.md", "Guía rápida para configuración, operación y restauración."),
]
for path, description in files:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.add_run(path).bold = True
    paragraph.add_run(" — " + description)

document.add_heading("5. Configuración paso a paso", level=1)
document.add_heading("5.1 Crear el secreto de GitHub", level=2)
for index, item in enumerate([
    "Entrar al repositorio en GitHub.",
    "Abrir Settings -> Secrets and variables -> Actions.",
    "Seleccionar New repository secret.",
    "Usar SUPABASE_DB_URL como nombre.",
    "Pegar la cadena de conexión PostgreSQL de Supabase con sslmode=require.",
    "Guardar el secreto sin publicarlo en archivos, issues o chats.",
], start=1):
    paragraph = document.add_paragraph(style="List Number")
    paragraph.add_run(item)
document.add_paragraph(
    "La cadena debe ser una URL PostgreSQL válida. El script acepta los modos "
    "sslmode=require, verify-ca o verify-full. Si falta TLS, el proceso se detiene."
)

document.add_heading("5.2 Ejecutar una prueba manual", level=2)
for item in [
    "Abrir la pestaña Actions del repositorio.",
    "Seleccionar el workflow Supabase database backup.",
    "Presionar Run workflow.",
    "Esperar a que el job termine con estado verde.",
    "Abrir la ejecución y descargar el artifact generado.",
]:
    add_bullet(document, item)

document.add_heading("5.3 Ejecución programada", level=2)
document.add_paragraph(
    "El workflow utiliza la expresión cron 0 5 * * *, que corresponde a las 05:00 UTC. "
    "GitHub puede retrasar algunos minutos los workflows programados durante periodos "
    "de alta demanda."
)

document.add_heading("6. Ejecución local en modo seguro", level=1)
document.add_paragraph(
    "El modo dry-run permite verificar la configuración sin conectarse a Supabase ni "
    "crear un archivo de backup."
)
add_code(document, '$env:SUPABASE_DB_URL = "postgresql://<usuario>:<clave>@<host>:5432/postgres?sslmode=require"\npowershell -NoProfile -ExecutionPolicy Bypass -File scripts/backup-supabase.ps1 -DryRun\nRemove-Item Env:SUPABASE_DB_URL')
document.add_paragraph(
    "Para ejecutar un backup real se necesita tener pg_dump instalado y una conexión "
    "válida. En GitHub Actions, pg_dump se instala automáticamente mediante el paquete "
    "postgresql-client."
)

document.add_heading("7. Restauración del backup", level=1)
document.add_paragraph(
    "La restauración debe realizarse primero sobre una base de datos de prueba. No se "
    "debe usar --clean contra producción sin una copia previa, aprobación y ventana de "
    "mantenimiento."
)
document.add_heading("7.1 Verificar el checksum", level=2)
add_code(document, "Get-FileHash .\\las-flores-public-YYYYMMDD-HHMMSS.dump -Algorithm SHA256\nGet-Content .\\las-flores-public-YYYYMMDD-HHMMSS.dump.sha256")
document.add_paragraph(
    "El hash calculado debe coincidir con el valor almacenado en el archivo .sha256. "
    "Si no coincide, no se debe intentar restaurar el dump."
)
document.add_heading("7.2 Restaurar en una base de destino", level=2)
add_code(document, '$env:RESTORE_DB_URL = "postgresql://<usuario>:<clave>@<host-destino>:5432/postgres?sslmode=require"\npg_restore --dbname=$env:RESTORE_DB_URL --clean --if-exists --no-owner --no-privileges .\\las-flores-public-YYYYMMDD-HHMMSS.dump\nRemove-Item Env:RESTORE_DB_URL')

document.add_heading("8. Seguridad y buenas prácticas", level=1)
for item in [
    "Mantener el repositorio privado, porque los dumps pueden contener datos personales.",
    "No guardar SUPABASE_DB_URL en .env, archivos YAML, scripts o documentación real.",
    "Rotar la credencial inmediatamente si se expone.",
    "Conservar los artifacts solo el tiempo necesario y controlar quién puede descargarlos.",
    "Probar periódicamente una restauración para verificar que el backup sea utilizable.",
    "Activar PITR en Supabase si el plan contratado lo permite.",
]:
    add_bullet(document, item)

document.add_heading("9. Validaciones realizadas", level=1)
validation = document.add_table(rows=1, cols=2)
validation.alignment = WD_TABLE_ALIGNMENT.CENTER
validation.style = "Light Shading Accent 1"
for cell, text in zip(validation.rows[0].cells, ["Validación", "Resultado"]):
    set_cell_text(cell, text, bold=True, color=(255, 255, 255))
    set_cell_shading(cell, "3E5C4E")
for check, result in [
    ("Smoke test del script de backup", "Correcto"),
    ("Suite existente del proyecto", "69 pruebas correctas"),
    ("Formato Prettier y whitespace", "Correcto"),
    ("Búsqueda de credenciales literales", "No se encontraron secretos"),
]:
    cells = validation.add_row().cells
    set_cell_text(cells[0], check)
    set_cell_text(cells[1], result)

document.add_heading("10. Lista de operación", level=1)
for item in [
    "[ ] Crear SUPABASE_DB_URL en GitHub Actions.",
    "[ ] Ejecutar el workflow manualmente.",
    "[ ] Confirmar que se generó el artifact.",
    "[ ] Descargar y verificar un checksum de prueba.",
    "[ ] Programar una restauración de prueba periódica.",
    "[ ] Definir un procedimiento separado para Storage y Auth.",
]:
    add_bullet(document, item)

document.add_paragraph()
closing = document.add_paragraph()
closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = closing.add_run("Fin del documento")
run.bold = True
run.font.color.rgb = RGBColor(62, 92, 78)

document.save(OUTPUT)
print(OUTPUT)